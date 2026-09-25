"""生成 Excel 统计表：总表 + 明细 + 计分规则。"""
from __future__ import annotations

from pathlib import Path

from . import rules


def _rc_text(r):
    out = [r["std"]] if r.get("std") else []
    if r.get("other"):
        out.append(f"执委({r['other']})")
    lines = ["、".join(out)] if out else []
    lines += [f"中层管理：{m}" for m in r.get("mid") or []]
    return "\n".join(lines)


def _mid_text(s, ov, year=None):
    mbb = rules.mid_by_block(s, ov, year)
    bl = [b for b in s["blocks"] if not year or b["year"] == year]
    return "\n".join(f"{b['year']} 中层管理：{x}" for b, m in zip(bl, mbb) for x in m)


def _roles_text(s, ov, year=None):
    _, rbb = rules.compute_stats(s, ov, year)
    bl = [b for b in s["blocks"] if not year or b["year"] == year]
    return "；".join(f"{b['year']}：{'、'.join(r)}" for b, r in zip(bl, rbb) if r)


def write_xlsx(students, out: Path, overrides: dict):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="DCE9F9")
    bold = Font(bold=True)

    def sheet(ws, header, rows, widths):
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = bold, head_fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    # 总表
    ws = wb.active
    ws.title = "总表"
    rows = []
    for s in students:
        st, _ = rules.compute_stats(s, overrides)
        rows.append([s["code"], s["club"], s["sid"], s["cn"], s["en"], s["cls"], len(set(s["years"])),
                     _roles_text(s, overrides), "、".join(s.get("special") or []), st["roles"], _mid_text(s, overrides), st["mid"], st["comm"], st["extComp"], st["intComp"], st["awards"],
                     st["extAct"], st["intAct"], st["team"], st["extSvc"], st["intSvc"], st["hours"], st["unsure"],
                     "；".join(s["issues"]), s["file"]])
    sheet(ws, ["学会代号", "学会", "学号", "姓名", "英文名", "班级", "年数", "职务（规范写法，历年）", "特别标记", "职务数", "中层管理（历年）", "中层管理数", "筹委",
               "校外比赛", "校内比赛", "获奖", "校外活动", "校内活动", "团内工作", "校外服务项", "校内服务项",
               "服务时数", "待确认", "资料问题", "文件"],
          rows, [8, 14, 8, 10, 22, 7, 6, 60, 14, 7, 40, 7, 6, 8, 8, 6, 8, 8, 8, 9, 9, 9, 7, 30, 40])
    for r in ws.iter_rows(min_row=2, min_col=8, max_col=11):
        for c in (r[0], r[3]):   # 职务、中层管理（每个一行）
            c.alignment = Alignment(wrap_text=True, vertical="top")

    # 明细
    ws2 = wb.create_sheet("明细")
    rows = []
    for s in students:
        for b in s["blocks"]:
            for k, arr in b["cats"].items():
                for i, t in enumerate(arr):
                    x = rules.item_state(s, b, k, i, overrides)
                    rows.append([s["code"], s["sid"], s["cn"], b["year"], b.get("cls") or "",
                                 (b.get("clubCode", "") + " " + b.get("clubName", "")).strip(),
                                 rules.CAT_LABEL.get(k, k), t,
                                 "计入" if x["inc"] else "不计", x["reason"] or "",
                                 "待确认" if x["unsure"] else "",
                                 "★" if k in ("extComp", "intComp") and b["aw"][k][i] else "",
                                 (b["hr"].get(k) or [None] * (i + 1))[i] if k in ("extSvc", "intSvc") else None,
                                 _rc_text(b["rc"][i]) if k == "role" and b["rc"][i] else "",
                                 (b.get("sp", {}).get(k) or [None] * (i + 1))[i] or ""])
    sheet(ws2, ["学会", "学号", "姓名", "年份", "班级", "该年学会", "类别", "内容", "计入", "不计原因", "待确认", "获奖", "小时", "职务规范写法", "特别标记"],
          rows, [7, 8, 10, 7, 7, 18, 12, 60, 6, 30, 8, 6, 7, 32, 12])
    grey = Font(color="999999")
    for r in ws2.iter_rows(min_row=2):
        if r[13].value and "\n" in str(r[13].value):
            r[13].alignment = Alignment(wrap_text=True, vertical="top")
        if r[8].value == "不计":
            for c in r:
                c.font = grey

    # 规则说明
    ws3 = wb.create_sheet("计分规则")
    for line in RULES_TEXT.strip().splitlines():
        ws3.append([line])
    ws3.column_dimensions["A"].width = 120
    wb.save(out)
    return out


RULES_TEXT = """
高三最高成就奖 · 计分规则（本工具自动套用）
1. 以班级为单位的内容、服务、活动、比赛不计（运动会写生、号码布设计、短片制作比赛也属班级活动）；学会的筹委、团内工作照算（例：合唱团筹办班级歌曲合唱比赛）。
2. 只计 A、C、D、E 类；B 类（体育、学术培训队）整年不计。
3. 同一年同一比赛写在两个学会时只计一次，保留与比赛相关的学会（例：舞蹈&口才 → 只算舞蹈）。
4. 比赛须代表本学会：与所属学会无关的个人比赛不计（例：舞蹈团学生的武术比赛）。按学会关键词自动判断，判断不了的标「待确认」，暂时计入。
执委/职务：按手册统一写法。标准职称：主席、秘书、事务（「事物」「总务」同）、财政（正/副）、查账、总学长（正/副）、助理总学长、执委、会员；其它职位写成「执委(XXX,XXX)」（计入其中的职位个数）。主席、秘书、事务、财政、总学长没写正/副的当作「正」。合唱团「音乐主席」、二十四节令鼓队的队长类与执委主席同级，写成「主席(正/副)」。童军等「分团」职位（队长类除外）算「会员」。「监督XX主席」「督导XXX主席」属于活动筹委，计入筹委。助理类、授课人类（含助教、小教练、教官）、队长类（含正/副队长、分团的总队长/小队长等，节令鼓队长除外）、监督/督导/顾问/教练/领队类、首席/组长/领养人类（如副弹拨首席、扬琴组长、扬琴组领养人、声部组长）、家族职位（如家族家长）、联课组别（如联课组别第六组）、校内服务负责人（「校内服务负责人：A、B」拆开，每项各算 1 个）从执委/会员分开，另计「中层管理」（中层管理数＝不同中层管理职位的个数，不算进职务数）。联课处工委等归入执委(…)。写在执委栏的活动筹委职位（如「《弈德杯》筹委主席」）和「校内服务负责人」移到筹委。感恩聚会不算学会活动，相关条目全部不计。运动会、田径赛的职位/筹委（如「运动会主席」「监督运动会服务主席」）算服务筹委团，计入筹委（班级名义的除外）；学会在运动会上的义卖不计，演出和服务照算；高三毕联会 / 高三年级组的职位和筹委（含编辑工委会、广告工委会、教师节工委会）属于年级，不计（相关服务、活动照算）。会员不计职务数，同年已有其它职位时不另列会员。
服务时数：优先采用学生自填的「总服务时数」；该年没填总数才用逐项相加；不计入的服务项会扣除其时数。
获奖：比赛条目含名次/奖项字眼（冠军、金奖、优秀奖、第X名、特优、佳作等）即计为获奖（自动识别）。
特别标记：联课处工委、文娱工委、XXX志工工委、国际交流筹委/负责人会特别标记出来（总表「特别标记」栏、明细「特别标记」栏），只做标记，不改变计分。
手动调整：在网页里点「计入 / 不计」，再用「导出手动调整」存成 JSON，下次运行工具时选择该文件即可套用。
"""

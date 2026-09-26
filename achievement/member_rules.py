"""读取 member_rules.json（执委 / 职务归类规则表），并提供「这条符不符合」的判断。

规则表的格式说明写在 member_rules.json 最上面的「说明」里。
查找顺序：
  1. 环境变量 ACHIEVEMENT_MEMBER_RULES 指定的文件
  2. config/member_rules.json（exe 所在文件夹 / 项目根目录下的 config/）——方便直接修改
  3. 打包进 exe 的那一份
文件改了（修改时间变了）会自动重新读取。
"""
from __future__ import annotations

import json
import re
from pathlib import Path

from .paths import config_candidates, find_config   # 规则文件放在 config/

FILE_NAME = "member_rules.json"
_WS = re.compile(r"[\s 　​-‏⁠-⁯﻿]+")


def _norm(s) -> str:
    return _WS.sub("", str(s or "")).lower()


def candidates(file_name=FILE_NAME, env_name="ACHIEVEMENT_MEMBER_RULES"):
    """规则文件的查找顺序：见 paths.config_candidates（config/ 优先）。"""
    return config_candidates(file_name, env_name)


def find_file(file_name, env_name) -> Path:
    return find_config(file_name, env_name)


def path() -> Path:
    return find_file(FILE_NAME, "ACHIEVEMENT_MEMBER_RULES")


class Rule:
    """一条规则：包含任一 / 正则（或）＋ 并且包含任一 ＋ 不包含 ＋ 学会或内容包含任一"""

    def __init__(self, d: dict):
        self.d = d
        self.name = d.get("名称") or d.get("标记") or ""
        self.any = [_norm(x) for x in d.get("包含任一") or []]
        self.rx = [re.compile(x, re.I) for x in d.get("正则") or []]
        self.also = [_norm(x) for x in d.get("并且包含任一") or []]
        self.none = [_norm(x) for x in d.get("不包含") or []]
        self.club_any = [_norm(x) for x in d.get("学会或内容包含任一") or []]

    def get(self, k, default=None):
        return self.d.get(k, default)

    def hit_first(self, text):
        """第一个出现的「包含任一」字眼（给「主席级」判断正/副用）"""
        t = _norm(text)
        for w in self.also + self.any:
            if w and w in t:
                return w
        return None

    def match(self, text, club="", skip_also=False) -> bool:
        t = _norm(text)
        if self.club_any and not any(w in _norm(club) + t for w in self.club_any):
            return False
        if self.any or self.rx:
            if not (any(w in t for w in self.any) or any(r.search(t) for r in self.rx)):
                return False
        if self.also and not skip_also and not any(w in t for w in self.also):
            return False
        if any(w in t for w in self.none):
            return False
        return True


SCOPES = {"执委栏和筹委栏": {"role", "comm"}, "全部栏目": None, "除筹委栏和团内工作以外": "except_comm_team"}
# 「适用栏目」也可以写成栏目名称的清单，例：["校内服务"]、["校内服务", "校外服务"]
COLS = {"执委栏": "role", "执委": "role", "筹委栏": "comm", "筹委": "comm", "校外比赛": "extComp", "校内比赛": "intComp",
        "校外活动": "extAct", "校内活动": "intAct", "校外服务": "extSvc", "校内服务": "intSvc", "团内工作": "team",
        "考章": "badge", "团内荣誉": "honor", "其他": "other"}


def _scope(v):
    if isinstance(v, list):
        return {COLS.get(x, x) for x in v}
    return SCOPES.get(v, None)


class MemberRules:
    def __init__(self, data: dict, source: Path | None = None):
        self.data, self.source = data, source
        self.exclude = [Rule(r) for r in data["一、不计"]["规则"]]
        self.to_comm = [Rule(r) for r in data["二、执委栏移到筹委栏"]["规则"]]
        self.comm_joiners = [w for w in data["二、执委栏移到筹委栏"].get("拆开连接词") or [] if w]
        self.to_role = [Rule(r) for r in data["三、筹委栏移到执委栏"]["规则"]]
        self.roles = [Rule(r) for r in data["四、职位归类"]["规则"]]
        t = data["四、职位归类"]["标准职称"]
        self.titles_side = list(t.get("有正副") or [])
        self.titles_plain = list(t.get("不分正副") or [])
        self.title_alias = {k: [x for x in v if x] for k, v in (t.get("同等写法") or {}).items()}
        self.not_title = [_norm(x) for x in t.get("不适用字眼") or []]
        self.special = [Rule(r) for r in data["五、特别标记"]["规则"]]
        self.mid_rules = [r for r in self.roles if r.get("归类") == "中层管理"]
        self.mid_words = [w for r in self.mid_rules for w in r.any]
        self.captain_words = [w for r in self.mid_rules if "队长" in r.any for w in r.any]
        self.split_words = [w for r in self.mid_rules if r.get("拆开") for w in r.any]
        self.grade_rules = [r for r in self.exclude if "毕联" in r.name or "年级" in r.name]

    # ---- 一、不计
    def exclusion(self, text, cat):
        for r in self.exclude:
            scope = _scope(r.get("适用栏目"))
            if scope == "except_comm_team" and cat in ("comm", "team"):
                continue
            if isinstance(scope, set) and cat not in scope:
                continue
            if r.match(text):
                return r.get("原因") or f"{r.name}，不计"
        return None

    def is_grade(self, text) -> bool:
        return any(r.match(text) for r in self.grade_rules)

    # ---- 二 / 三
    def moves_to_comm(self, text) -> bool:
        return any(r.match(text) for r in self.to_comm)

    def moves_to_role(self, text) -> bool:
        return any(r.match(text) for r in self.to_role) and not self.moves_to_comm(text)

    # ---- 四
    def role_rule(self, text, club=""):
        for r in self.roles:
            if r.match(text, club):
                return r
        return None

    def has_mid(self, text) -> bool:
        t = _norm(text)
        return any(w in t for w in self.mid_words)

    # ---- 五
    def special_label(self, text, cat=None):
        for r in self.special:
            if r.match(text, skip_also=bool(r.get("筹委栏都算")) and cat == "comm"):
                return r.get("标记")
        return None


_cache = {"mtime": None, "path": None, "rules": None}


def get() -> MemberRules:
    p = path()
    m = p.stat().st_mtime_ns
    if _cache["rules"] is None or _cache["path"] != p or _cache["mtime"] != m:
        try:
            with open(p, encoding="utf-8") as f:
                data = json.load(f)
        except json.JSONDecodeError as e:
            raise ValueError(f"member_rules.json 格式有误（第 {e.lineno} 行第 {e.colno} 个字附近）：{e.msg}。"
                             "常见原因：少了逗号、多了最后一个逗号、用了中文引号。") from None
        _cache.update(rules=MemberRules(data, p), path=p, mtime=m)
    return _cache["rules"]


def fingerprint() -> str:
    """规则文件的指纹（改了规则 → 网页自动重算）"""
    try:
        p = path()
        st = p.stat()
        return f"{p}|{st.st_mtime_ns}|{st.st_size}"
    except OSError:
        return ""


# ---------------------------------------------------------------- 已知职位一览（参考用，另存成 Excel）


def _batches(students):
    """students 可以是学生清单，或 {届别: 学生清单}（Result 里有年份子文件夹时）"""
    return students if isinstance(students, dict) else {"": students}


def known_roles(students):
    """把 Result 里所有执委栏（含移栏）的职位，连同系统目前的归类整理成清单。
    students 是 {届别: 学生清单} 时，另外记下每一届出现的次数（row["届别"]）。"""
    from collections import Counter, defaultdict
    from . import rules
    R = get()
    stat = {}
    clubs = defaultdict(set)
    cnt = Counter()
    by_batch = defaultdict(Counter)
    batch = ""

    def add(text, info, club):
        key = rules.norm(text).rstrip("。.;；,，")
        if not key:
            return
        stat.setdefault(key, {"_原文": str(text).strip().rstrip("。.;；,，"), **info})
        cnt[key] += 1
        by_batch[key][batch] += 1
        if club:
            clubs[key].add(club)

    for batch, s in ((b, s) for b, ss in _batches(students).items() for s in ss):
        for b in s["blocks"]:
            club = b.get("clubCode") or ""
            ex = (b.get("ex") or {})
            for i, t in enumerate(b["cats"].get("role", [])):
                e = b.get("exBlock") or (ex.get("role") or [None] * (i + 1))[i]
                r = b["rc"][i] or {}
                rule = R.role_rule(rules.norm(t), b.get("clubName") or "")
                if e:
                    info = {"归类": "不计", "依据": e}
                elif r.get("mid"):
                    info = {"归类": "中层管理", "写法": "、".join(r["mid"]), "依据": rule.name if rule else "中层管理"}
                    if r.get("std") or r.get("other"):
                        info["同一条的其它职位"] = r.get("std") or f"执委({r['other']})"
                elif r.get("std") == "会员":
                    info = {"归类": "会员", "依据": rule.name if rule else "会员"}
                elif r.get("std", "").startswith("主席(") and rule and rule.get("归类") == "主席级":
                    info = {"归类": "主席级", "写法": r["std"], "依据": rule.name}
                elif r.get("std"):
                    info = {"归类": "执委", "写法": r["std"], "依据": rule.name if rule else "标准职称"}
                elif r.get("other"):
                    info = {"归类": "执委", "写法": f"执委({r['other']})", "依据": "其它职位"}
                else:
                    continue
                if b.get("mv", {}).get("role") and b["mv"]["role"][i]:
                    info["备注"] = "原写在筹委栏"
                add(t, info, club)
            moved = set(b.get("movedComm") or [])
            for i, t in enumerate(b["cats"].get("comm", [])):
                if t not in moved:
                    continue
                e = b.get("exBlock") or (ex.get("comm") or [None] * (i + 1))[i]
                why = next((x.name for x in R.to_comm if x.match(t)), "")
                add(t, {"归类": "不计" if e else "筹委", "依据": e or f"执委栏移到筹委栏：{why}", "备注": "原写在执委栏"}, club)
    order = {"执委": 0, "主席级": 1, "中层管理": 2, "会员": 3, "筹委": 4, "不计": 5}
    out = []
    for k, info in stat.items():
        info = dict(info)
        row = {"原文": info.pop("_原文", k), **info, "次数": cnt[k], "届别": dict(by_batch[k]), "学会": sorted(clubs[k])}
        out.append(row)
    out.sort(key=lambda r: (order.get(r["归类"], 9), r.get("依据", ""), -r["次数"], r["原文"]))
    return out


def write_known_xlsx(students, out, date_text=""):
    """把 Result 里所有执委栏职位及目前的归类，写成 Excel（只供参考，不影响计算）。out 可以是路径或 BytesIO。"""
    from collections import Counter
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    rows = known_roles(students)
    batches = [b for b in _batches(students) if b]
    wb = Workbook()
    ws = wb.active
    ws.title = "职位一览"
    head = ["原文", "归类", "写法", "依据", "同一条的其它职位", "备注", "次数", *[f"{b}届" for b in batches], "学会"]
    ws.append(head)
    for r in rows:
        ws.append([r.get("原文"), r.get("归类"), r.get("写法", ""), r.get("依据", ""), r.get("同一条的其它职位", ""),
                   r.get("备注", ""), r.get("次数"), *[r["届别"].get(b) or None for b in batches],
                   "、".join(r.get("学会") or [])])
    fills = {"执委": "DCE9F9", "主席级": "C9DAF8", "中层管理": "FFF2CC", "会员": "EDEDED", "筹委": "D9EAD3", "不计": "F4CCCC"}
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="BDD7EE")
    for row in ws.iter_rows(min_row=2):
        f = fills.get(row[1].value)
        if f:
            row[1].fill = PatternFill("solid", fgColor=f)
        row[0].alignment = row[2].alignment = Alignment(wrap_text=True, vertical="top")
    from openpyxl.utils import get_column_letter
    for i, w in enumerate([46, 10, 40, 34, 22, 14, 7, *[8] * len(batches), 22], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws2 = wb.create_sheet("说明")
    summary = Counter(r["归类"] for r in rows)
    for line in [
        "已知职位一览（只供参考）",
        f"产生时间：{date_text}",
        "Result 里出现过的所有执委栏职位（包括原写在执委栏、被移到筹委的），以及系统目前把它算成什么。",
        *(["Result 里有好几届（年份子文件夹）：「次数」是全部加起来，右边「XXXX届」是各届分别出现的次数。"] if batches else []),
        "改这份 Excel 不会改变计算；要改归类，请改 member_rules.json 的「一」到「四」，再重新产生这份清单。",
        "重新产生：网页右上角「下载职位一览」，或 python app.py --list-roles",
        "",
        "各归类的职位数：",
        *[f"  {k}：{v}" for k, v in summary.items()],
    ]:
        ws2.append([line])
    ws2.column_dimensions["A"].width = 100
    wb.save(out)
    return len(rows), summary

"""读取 award.json（比赛条目算不算获奖的规则表），并判断一条比赛条目是否获奖。

规则表的格式说明写在 award.json 最上面的「说明」里；查找位置和 member_rules.json 一样：
  1. 环境变量 ACHIEVEMENT_AWARD_RULES 指定的文件
  2. config/award.json（exe 所在文件夹 / 项目根目录下的 config/）
  3. 打包进 exe 的那一份
文件改了会自动重新读取。
"""
from __future__ import annotations

import json
import re

from .member_rules import _norm
from .paths import find_config as find_file

FILE_NAME = "award.json"
_PART_SPLIT = re.compile(r"——|--|—|–|－|：|:|\s-\s|(?<=\S)-(?=\S*奖)|-\s+(?=\S*奖)")


def path():
    return find_file(FILE_NAME, "ACHIEVEMENT_AWARD_RULES")


class AwardRules:
    def __init__(self, d: dict):
        a, b, c, e = d["一、不算获奖（例外）"], d["二、出现即算获奖"], d["三、比赛名称后面写的名次 / 奖项"], d["四、写明「获得 / 荣获 / 赢得」"]
        self.no_any = [_norm(x) for x in a.get("包含任一") or []]
        self.no_rx = [re.compile(x, re.I) for x in a.get("正则") or []]
        self.any = [_norm(x) for x in b.get("包含任一") or []]
        self.rx = [re.compile(x, re.I) for x in b.get("正则") or []]
        self.tail_any = [_norm(x) for x in c.get("后段包含任一") or []]
        self.tail_rx = [re.compile(x, re.I) for x in c.get("后段正则") or []]
        self.got_rx = [re.compile(x, re.I) for x in e.get("正则") or []]
        f = d.get("五、比赛是否代表本学会") or {}
        comp = lambda v: ([_norm(x) for x in v.get("关键词") or []], [re.compile(x, re.I) for x in v.get("正则") or []])
        self.club_kw = {code: comp(v) for code, v in (f.get("各学会关键词") or {}).items()}
        self.club_name = {code: v.get("名称", "") for code, v in (f.get("各学会关键词") or {}).items()}
        self.other_kw = {name: comp(v) for name, v in (f.get("其它类别关键词") or {}).items()}
        self.all_kw = comp(f.get("所有学会都算") or {})

    @staticmethod
    def _hit(kw, text):
        words, rxs = kw
        t = _norm(text)
        return any(w and w in t for w in words) or any(r.search(str(text)) for r in rxs)

    def own(self, code, text) -> bool:
        """条目里有没有出现这个学会的关键词"""
        kw = self.club_kw.get(code or "")
        return bool(kw and self._hit(kw, text))

    def comp_judge(self, text, code):
        """比赛是否代表本学会：None=计入；"unsure"=待确认；其它字串=不计原因"""
        if self.own(code, text) or self._hit(self.all_kw, text):
            return None
        for lab, kw in self.other_kw.items():
            if self._hit(kw, text):
                return f"非代表本学会（{lab}类比赛）"
        for c, kw in self.club_kw.items():
            if c != code and self._hit(kw, text):
                return f"非代表本学会（看似 {c} 类比赛）"
        return "unsure" if (code or "") in self.club_kw else None

    def why(self, text):
        """算获奖 → 返回依据（字串）；不算 → None"""
        s = str(text or "")
        t = _norm(s)
        if any(w and w in t for w in self.no_any) or any(r.search(s) for r in self.no_rx):
            return None
        for w in self.any:
            if w and w in t:
                return f"出现「{w}」"
        for r in self.rx:
            m = r.search(s)
            if m:
                return f"出现「{m.group(0)}」"
        parts = _PART_SPLIT.split(s)
        if len(parts) > 1:
            tail = "".join(parts[1:])
            tn = _norm(tail)
            for w in self.tail_any:
                if w and w in tn:
                    return f"后段写了「{w}」"
            for r in self.tail_rx:
                m = r.search(tail)
                if m:
                    return f"后段写了「{m.group(0)}」"
        for r in self.got_rx:
            m = r.search(s)
            if m:
                return f"写明「{m.group(0)}」"
        return None


_cache = {"mtime": None, "path": None, "rules": None}


def get() -> AwardRules:
    p = path()
    m = p.stat().st_mtime_ns
    if _cache["rules"] is None or _cache["path"] != p or _cache["mtime"] != m:
        try:
            with open(p, encoding="utf-8") as f:
                data = json.load(f)
        except json.JSONDecodeError as e:
            raise ValueError(f"award.json 格式有误（第 {e.lineno} 行第 {e.colno} 个字附近）：{e.msg}。"
                             "常见原因：少了逗号、多了最后一个逗号、用了中文引号。") from None
        _cache.update(rules=AwardRules(data), path=p, mtime=m)
    return _cache["rules"]


def fingerprint() -> str:
    try:
        p = path()
        st = p.stat()
        return f"{p}|{st.st_mtime_ns}|{st.st_size}"
    except OSError:
        return ""


def write_known_xlsx(students, out, date_text=""):
    """所有比赛条目 + 目前算不算获奖（只供参考）。"""
    from collections import Counter, defaultdict
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from . import rules as RL
    from .member_rules import _batches
    A = get()
    stat, cnt, clubs, by_batch = {}, Counter(), defaultdict(set), defaultdict(Counter)
    batches = [b for b in _batches(students) if b]
    for batch, s in ((bn, s) for bn, ss in _batches(students).items() for s in ss):
        for b in s["blocks"]:
            for k, lab in (("extComp", "校外比赛"), ("intComp", "校内比赛")):
                for i, t in enumerate(b["cats"].get(k, [])):
                    key = RL.norm(t).rstrip("。.;；,，")
                    if not key:
                        continue
                    if key not in stat:
                        w = A.why(t)
                        e = b.get("exBlock") or (b.get("ex", {}).get(k) or [None] * (i + 1))[i]
                        stat[key] = {"原文": str(t).strip(), "类别": lab, "获奖": "★ 获奖" if w else "—", "依据": w or "", "计入": "不计：" + e if e else "计入"}
                    cnt[key] += 1
                    by_batch[key][batch] += 1
                    if b.get("clubCode"):
                        clubs[key].add(b["clubCode"])
    rows = sorted(stat.items(), key=lambda kv: (kv[1]["获奖"] != "★ 获奖", kv[1]["依据"], -cnt[kv[0]], kv[0]))
    wb = Workbook()
    ws = wb.active
    ws.title = "获奖一览"
    ws.append(["原文", "类别", "获奖", "依据", "这条是否计入", "次数", *[f"{b}届" for b in batches], "学会"])
    for k, v in rows:
        ws.append([v["原文"], v["类别"], v["获奖"], v["依据"], v["计入"], cnt[k],
                   *[by_batch[k].get(b) or None for b in batches], "、".join(sorted(clubs[k]))])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="BDD7EE")
    gold, grey = PatternFill("solid", fgColor="FFF2CC"), Font(color="999999")
    for row in ws.iter_rows(min_row=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
        if row[2].value == "★ 获奖":
            row[2].fill = gold
        if str(row[4].value).startswith("不计"):
            for c in row:
                c.font = grey
    from openpyxl.utils import get_column_letter
    for i, w in enumerate([60, 9, 9, 22, 30, 7, *[8] * len(batches), 18], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    summary = Counter(v["获奖"] for _, v in rows)
    ws2 = wb.create_sheet("说明")
    for line in ["已知获奖一览（只供参考）", f"产生时间：{date_text}",
                 "Result 里出现过的所有比赛条目（校外 / 校内），以及系统目前算不算获奖、依据哪条规则。",
                 "改这份 Excel 不会改变计算；要改判断，请改 award.json，再重新产生这份清单。",
                 "「这条是否计入」是另外的规则（班级、非代表本学会、B 类等），和获奖与否无关。",
                 "重新产生：网页右上角「下载获奖一览」，或 python app.py --list-awards", "",
                 f"获奖：{summary.get('★ 获奖', 0)} 条　　未获奖：{summary.get('—', 0)} 条（不同写法各算一条）"]:
        ws2.append([line])
    ws2.column_dimensions["A"].width = 100
    wb.save(out)
    return len(rows), summary
